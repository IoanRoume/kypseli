# Copyright 2026 Ioannis Roumeliotis
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#
from file_organizer.extractors.base import BaseExtractor
from file_organizer.core.models import FileInfo, ExtractedContent
import pypdf
from docx import Document
from pptx import Presentation
from striprtf.striprtf import rtf_to_text
import ebooklib
from ebooklib import epub
from bs4 import BeautifulSoup
import warnings

class DocumentExtractor(BaseExtractor):
    supported_extensions = [
        ".pdf",
        ".docx",
        ".pptx",
        ".rtf",
        ".epub"
    ]


    def extract(self, file_info: FileInfo, max_chars: int = 7000) -> ExtractedContent:
        file_path = file_info.path
        ext = file_path.suffix.lower()

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        text_content = ""
        
        try:
            if ext == '.pdf':
                reader = pypdf.PdfReader(file_path)
                extracted_text = []
                char_count = 0
                
                for page in reader.pages:
                    if reader.is_encrypted:
                        try:
                            reader.decrypt("")
                        except:
                            return ExtractedContent(file_info, "Error: PDF is encrypted", "failed")

                    text = page.extract_text() or ""
                    extracted_text.append(text)
                    char_count += len(text)
                    if char_count >= max_chars:
                        break
                text_content = "\n".join(extracted_text)

            elif ext == '.docx':
                doc = Document(file_path)
                text_content = "\n".join([para.text for para in doc.paragraphs])

            elif ext == '.pptx':
                prs = Presentation(file_path)
                text_parts = []
                
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text_frame") and shape.text_frame:
                            for paragraph in shape.text_frame.paragraphs:
                                text_parts.append(paragraph.text)
                text_content = "\n".join(text_parts)

            elif ext == '.rtf':
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    rtf_content = f.read()
                    text_content = rtf_to_text(rtf_content)

            elif ext == '.epub':
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore")
                    book = epub.read_epub(file_path)
                    
                chapters = []
                for item in book.get_items():
                    if item.get_type() == ebooklib.ITEM_DOCUMENT:
                        soup = BeautifulSoup(item.get_content(), 'html.parser')
                        chapters.append(soup.get_text())
                text_content = "\n".join(chapters)

            else:
                return ExtractedContent(file_info, "Unsupported document format", "failed")

        except Exception as e:
            return ExtractedContent(
                file_info=file_info,
                content=f"Error reading document: {str(e)}",
                extraction_method="failed"
            )

        content = text_content[:max_chars]

        return ExtractedContent(
            file_info=file_info,
            content=content,
            extraction_method="document_parse"
        )
